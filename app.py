from flask import Flask, redirect, render_template, request, url_for, flash, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import date, datetime
from sqlalchemy import extract
import sqlite3
from pathlib import Path
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///expenses.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SECRET_KEY"] = "change-this-to-a-long-random-secret"
db = SQLAlchemy(app)

DEFAULT_CATEGORIES = [
    "Food",
    "Transport",
    "Rent",
    "Utilities",
    "Entertainment",
    "Healthcare",
    "Other",
]

PAYMENT_METHODS = [
    "Cash",
    "UPI",
    "Debit Card",
    "Credit Card",
    "Net Banking",
    "Online Payment",
    "Other",
]


class Expense(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    description = db.Column(db.String(120), nullable=False)
    amount = db.Column(db.Float, nullable=False)
    category = db.Column(db.String(50), nullable=False)
    date = db.Column(db.Date, nullable=False, default=date.today)
    payment_method = db.Column(
        db.String(30), nullable=False, default="Cash"
    )


class UserProfile(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    display_name = db.Column(
        db.String(80), nullable=False, default="My Profile"
    )
    email = db.Column(db.String(120), nullable=False, default="")
    currency = db.Column(db.String(8), nullable=False, default="INR")


class AppCategory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(50), unique=True, nullable=False)


def ensure_payment_method_column():
    """Add payment_method to an existing SQLite table without deleting data."""
    with app.app_context():
        db.create_all()

        database_path = Path(app.instance_path) / "expenses.db"
        database_path.parent.mkdir(parents=True, exist_ok=True)

        connection = sqlite3.connect(database_path)
        try:
            columns = connection.execute(
                "PRAGMA table_info(expense)"
            ).fetchall()
            column_names = [column[1] for column in columns]

            if "payment_method" not in column_names:
                connection.execute(
                    "ALTER TABLE expense "
                    "ADD COLUMN payment_method VARCHAR(30) "
                    "NOT NULL DEFAULT 'Cash'"
                )
                connection.commit()
        finally:
            connection.close()

        if UserProfile.query.first() is None:
            db.session.add(
                UserProfile(
                    display_name="My Profile",
                    email="",
                    currency="INR",
                )
            )

        for category_name in DEFAULT_CATEGORIES:
            if not AppCategory.query.filter_by(name=category_name).first():
                db.session.add(AppCategory(name=category_name))

        db.session.commit()


ensure_payment_method_column()


def get_categories():
    return [
        category.name
        for category in AppCategory.query.order_by(AppCategory.name).all()
    ]


def get_profile():
    return UserProfile.query.first()


def parse_date(value):
    return datetime.strptime(value, "%Y-%m-%d").date()


def filtered_expenses(args):
    query = Expense.query
    specific_date = (args.get("specific_date") or "").strip()
    start_date = (args.get("start_date") or "").strip()
    end_date = (args.get("end_date") or "").strip()
    selected_month = (args.get("month") or "").strip()
    selected_year = (args.get("year") or "").strip()
    category = (args.get("category") or "").strip()
    search = (args.get("search") or "").strip()

    if specific_date:
        try:
            query = query.filter(Expense.date == parse_date(specific_date))
        except ValueError:
            flash("Please enter a valid specific date.", "error")
    else:
        if start_date:
            try:
                query = query.filter(Expense.date >= parse_date(start_date))
            except ValueError:
                flash("Invalid start date.", "error")

        if end_date:
            try:
                query = query.filter(Expense.date <= parse_date(end_date))
            except ValueError:
                flash("Invalid end date.", "error")

    if selected_month:
        try:
            month_number = int(selected_month)
            if not 1 <= month_number <= 12:
                raise ValueError
            query = query.filter(extract("month", Expense.date) == month_number)
        except ValueError:
            flash("Please select a valid month.", "error")
            selected_month = ""

    if selected_year:
        try:
            year_number = int(selected_year)
            if not 2000 <= year_number <= 2100:
                raise ValueError
            query = query.filter(extract("year", Expense.date) == year_number)
        except ValueError:
            flash("Please select a valid year.", "error")
            selected_year = ""

    if category:
        if category in get_categories():
            query = query.filter(Expense.category == category)
        else:
            flash("Please select a valid category.", "error")
            category = ""

    if search:
        query = query.filter(Expense.description.ilike(f"%{search}%"))

    expenses = query.order_by(Expense.date.desc(), Expense.id.desc()).all()

    return {
        "expenses": expenses,
        "total": round(sum(expense.amount for expense in expenses), 2),
        "specific_date": specific_date,
        "start_date": start_date,
        "end_date": end_date,
        "selected_month": selected_month,
        "selected_year": selected_year,
        "category": category,
        "selected_category": category,
        "search": search,
        "categories": get_categories(),
        "today": date.today().isoformat(),
        "profile": get_profile(),
    }


@app.context_processor
def shared_template_values():
    return {
        "site_profile": get_profile(),
        "site_categories": get_categories(),
    }


@app.route("/")
def home():
    return render_template("index.html", **filtered_expenses(request.args))


@app.route("/transactions")
def transactions():
    return render_template(
        "transactions.html", **filtered_expenses(request.args)
    )


@app.route("/analytics")
def analytics():
    data = filtered_expenses(request.args)
    category_totals = {}
    month_totals = {}

    for expense in data["expenses"]:
        category_totals[expense.category] = (
            category_totals.get(expense.category, 0) + expense.amount
        )
        month_key = expense.date.strftime("%Y-%m")
        month_totals[month_key] = (
            month_totals.get(month_key, 0) + expense.amount
        )

    data["category_labels"] = list(category_totals.keys())
    data["category_values"] = [
        round(value, 2) for value in category_totals.values()
    ]
    data["month_labels"] = sorted(month_totals.keys())
    data["month_values"] = [
        round(month_totals[key], 2) for key in sorted(month_totals)
    ]
    data["average"] = (
        round(data["total"] / len(data["expenses"]), 2)
        if data["expenses"]
        else 0
    )
    data["highest"] = max(
        data["expenses"], key=lambda expense: expense.amount, default=None
    )

    return render_template("analytics.html", **data)


@app.route("/add-expense")
def add_expense_page():
    return render_template(
        "add_expense.html",
        categories=get_categories(),
        today=date.today().isoformat(),
        profile=get_profile(),
        payment_methods=PAYMENT_METHODS,
    )


@app.route("/add", methods=["POST"])
def add():
    description = (request.form.get("description") or "").strip()
    raw_amount = (request.form.get("amount") or "").strip()
    category = (request.form.get("category") or "").strip()
    date_string = (request.form.get("date") or "").strip()
    payment_method = (request.form.get("payment_method") or "Cash").strip()

    if not description or len(description) > 120 or not raw_amount or not category:
        flash("Please enter a description, amount, and category.", "error")
        return redirect(request.referrer or url_for("add_expense_page"))

    if category not in get_categories():
        flash("Please select a valid category.", "error")
        return redirect(request.referrer or url_for("add_expense_page"))

    if payment_method not in PAYMENT_METHODS:
        flash("Please select a valid payment method.", "error")
        return redirect(request.referrer or url_for("add_expense_page"))

    try:
        amount = float(raw_amount)
        if amount <= 0:
            raise ValueError
    except ValueError:
        flash("Amount must be a positive number.", "error")
        return redirect(request.referrer or url_for("add_expense_page"))

    try:
        expense_date = parse_date(date_string) if date_string else date.today()
    except ValueError:
        flash("Please enter a valid date.", "error")
        return redirect(request.referrer or url_for("add_expense_page"))

    expense = Expense(
        description=description,
        amount=amount,
        category=category,
        date=expense_date,
        payment_method=payment_method,
    )
    db.session.add(expense)
    db.session.commit()

    flash("Expense added successfully!", "success")
    return redirect(request.referrer or url_for("home"))


@app.route("/edit/<int:expense_id>", methods=["GET", "POST"])
def edit(expense_id):
    expense = Expense.query.get_or_404(expense_id)

    if request.method == "POST":
        description = (request.form.get("description") or "").strip()
        raw_amount = (request.form.get("amount") or "").strip()
        category = (request.form.get("category") or "").strip()
        date_string = (request.form.get("date") or "").strip()
        payment_method = (
            request.form.get("payment_method") or expense.payment_method or "Cash"
        ).strip()

        try:
            amount = float(raw_amount)
            expense_date = parse_date(date_string)
            if (
                not description
                or len(description) > 120
                or amount <= 0
                or category not in get_categories()
                or payment_method not in PAYMENT_METHODS
            ):
                raise ValueError
        except ValueError:
            flash(
                "Check the description, amount, category, date, and payment method.",
                "error",
            )
            return render_template(
                "edit.html",
                expense=expense,
                categories=get_categories(),
                profile=get_profile(),
                payment_methods=PAYMENT_METHODS,
            )

        expense.description = description
        expense.amount = amount
        expense.category = category
        expense.date = expense_date
        expense.payment_method = payment_method

        db.session.commit()
        flash("Expense updated successfully!", "success")
        return redirect(url_for("transactions"))

    return render_template(
        "edit.html",
        expense=expense,
        categories=get_categories(),
        profile=get_profile(),
        payment_methods=PAYMENT_METHODS,
    )


@app.route("/delete/<int:expense_id>", methods=["POST"])
def delete(expense_id):
    expense = Expense.query.get_or_404(expense_id)
    db.session.delete(expense)
    db.session.commit()
    flash("Expense deleted successfully!", "success")
    return redirect(request.referrer or url_for("transactions"))


@app.route("/delete-multiple", methods=["POST"])
def delete_multiple():
    selected_ids = request.form.getlist("expense_ids")
    if not selected_ids:
        flash("Please select at least one expense.", "error")
        return redirect(request.referrer or url_for("transactions"))

    expenses = Expense.query.filter(Expense.id.in_(selected_ids)).all()
    count = len(expenses)

    for expense in expenses:
        db.session.delete(expense)

    db.session.commit()
    flash(f"{count} expense(s) deleted successfully.", "success")
    return redirect(request.referrer or url_for("transactions"))


@app.route("/profile", methods=["GET", "POST"])
def profile():
    user_profile = get_profile()

    if request.method == "POST":
        name = (request.form.get("display_name") or "").strip()
        email = (request.form.get("email") or "").strip()

        if not name:
            flash("Display name is required.", "error")
        elif len(name) > 80 or len(email) > 120:
            flash("Please use shorter profile details.", "error")
        else:
            user_profile.display_name = name
            user_profile.email = email
            db.session.commit()
            flash("Profile saved successfully.", "success")
            return redirect(url_for("profile"))

    return render_template("profile.html", profile=user_profile)


@app.route("/settings", methods=["GET", "POST"])
def settings():
    user_profile = get_profile()

    if request.method == "POST":
        action = request.form.get("action", "")

        if action == "currency":
            currency = request.form.get("currency", "INR")
            if currency not in {"INR", "USD", "GBP", "EUR"}:
                flash("Invalid currency selection.", "error")
            else:
                user_profile.currency = currency
                db.session.commit()
                flash(
                    "Currency preference saved. Existing display labels use ₹ "
                    "in this starter version.",
                    "success",
                )

        elif action == "add_category":
            name = (request.form.get("category_name") or "").strip()
            if not name or len(name) > 50:
                flash("Enter a category name up to 50 characters.", "error")
            elif AppCategory.query.filter_by(name=name).first():
                flash("That category already exists.", "error")
            else:
                db.session.add(AppCategory(name=name))
                db.session.commit()
                flash("Category added.", "success")

        elif action == "delete_category":
            name = (request.form.get("category_name") or "").strip()
            if Expense.query.filter_by(category=name).first():
                flash(
                    "This category is used by existing expenses, so it cannot be deleted yet.",
                    "error",
                )
            else:
                category = AppCategory.query.filter_by(name=name).first()
                if category:
                    db.session.delete(category)
                    db.session.commit()
                    flash("Category deleted.", "success")

        return redirect(url_for("settings"))

    return render_template(
        "settings.html",
        profile=user_profile,
        categories=get_categories(),
    )


@app.route("/export/csv")
def export_csv():
    """Download expenses matching the current filters as a CSV file."""
    expenses = filtered_expenses(request.args)["expenses"]

    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(
        ["ID", "Date", "Description", "Category", "Payment Method", "Amount (INR)"]
    )

    for expense in expenses:
        writer.writerow(
            [
                expense.id,
                expense.date.isoformat(),
                expense.description,
                expense.category,
                expense.payment_method or "Cash",
                round(expense.amount, 2),
            ]
        )

    memory_file = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    memory_file.seek(0)

    return send_file(
        memory_file,
        mimetype="text/csv; charset=utf-8",
        as_attachment=True,
        download_name="expenseease_transactions.csv",
    )


@app.route("/export/excel")
def export_excel():
    """Download expenses matching the current filters as an Excel workbook."""
    expenses = filtered_expenses(request.args)["expenses"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"

    headers = [
        "ID",
        "Date",
        "Description",
        "Category",
        "Payment Method",
        "Amount (INR)",
    ]
    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for expense in expenses:
        sheet.append(
            [
                expense.id,
                expense.date.isoformat(),
                expense.description,
                expense.category,
                expense.payment_method or "Cash",
                round(expense.amount, 2),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    column_widths = {
        "A": 10,
        "B": 14,
        "C": 32,
        "D": 20,
        "E": 22,
        "F": 18,
    }
    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    for cell in sheet["F"][1:]:
        cell.number_format = '#,##0.00'

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="expenseease_transactions.xlsx",
    )



if __name__ == "__main__":
    app.run(debug=True, port=4848)
